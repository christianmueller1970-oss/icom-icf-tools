#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 Chris, HB9HJI – https://www.funkwelt.net
# Teil von icom-icf-tools. Ohne CHIRP lauffähig (openpyxl nur für Excel-Master).
"""
make_device_lists.py – erzeugt aus der Master-Repeaterliste die Import-CSV je Gerät.

Aufruf:  python3 make_device_lists.py master.xlsx|master.csv out/ [JJJJMMTT]

Master = Icom-Repeaterlisten-Format (17 Spalten wie ID-52-Export, Dezimalpunkt),
als Excel (Blatt «Repeater») oder als CSV (Komma oder Semikolon).

Ergebnis (Dateiname wie vom Gerät selbst vergeben, NN = 10, damit nichts überschrieben wird):
  out/ID-52E/Rpt<datum>_10.csv    Komma, Dezimalpunkt, 17 Spalten, 2 m + 70 cm
  out/IC-705/Rpt<datum>_10.csv    Semikolon, Dezimalkomma, 17 Spalten, KW/6 m/2 m/70 cm
  out/ID-5100/Rpt<datum>_10.csv   Semikolon, Dezimalkomma, 17 Spalten, 2 m + 70 cm
  out/IC-9700/Rpt<datum>_10.csv   Semikolon, Dezimalkomma, 18 Spalten (+ Gateway IP Address), 2 m/70 cm/23 cm
  out/IC-9700/Rpt<datum>_11.csv   dito, nur DV/DD (Rückfallebene, falls der IC-9700 FM-Einträge ablehnt)

Formate = exakt die Header der Geräte-Exporte (Firmware-Stand 09/2026), siehe docs/Icom-CSV-Formate.md.
Gruppen werden je Gerät lückenlos neu nummeriert (1, 2, 3 …), leere Gruppen fallen weg.
"""
import csv, os, sys, datetime
from collections import Counter, OrderedDict

COLS17 = ["Group No", "Group Name", "Name", "Sub Name", "Repeater Call Sign", "Gateway Call Sign",
          "Frequency", "Dup", "Offset", "Mode", "TONE", "Repeater Tone", "RPT1USE", "Position",
          "Latitude", "Longitude", "UTC Offset"]
COLS18 = COLS17[:6] + ["Gateway IP Address"] + COLS17[6:]
NUM = ("Frequency", "Offset", "Latitude", "Longitude", "Repeater Tone")

BANDS = {"HF": (28.0, 29.7), "6m": (50.0, 52.0), "2m": (144.0, 146.0), "70cm": (430.0, 440.0),
         "23cm": (1240.0, 1300.0)}
DEVICES = OrderedDict([
    ("ID-52E",  dict(sep=",", dec=".", cols=COLS17, bands=("2m", "70cm"), limit=2500)),
    ("IC-705",  dict(sep=";", dec=",", cols=COLS17, bands=("HF", "6m", "2m", "70cm"), limit=2500)),
    ("ID-5100", dict(sep=";", dec=",", cols=COLS17, bands=("2m", "70cm"), limit=1500)),
    ("IC-9700", dict(sep=";", dec=",", cols=COLS18, bands=("2m", "70cm", "23cm"), limit=2500)),
])


def band(f):
    for b, (lo, hi) in BANDS.items():
        if lo <= f <= hi:
            return b
    return None


def read_master(path):
    if path.lower().endswith(".csv"):
        t = open(path, encoding="utf-8-sig").read()
        d = ";" if t.split("\n")[0].count(";") > 3 else ","
        rows = list(csv.DictReader(t.splitlines(), delimiter=d))
        assert list(rows[0].keys()) == COLS17, "Master-Header weicht ab"
        for r in rows:
            for c in ("Frequency", "Offset", "Latitude", "Longitude", "Repeater Tone"):
                r[c] = (r[c] or "").replace(",", ".")
        return [{c: (r[c] or "").strip() for c in COLS17} for r in rows]
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True)["Repeater"]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h) for h in next(it)]
    assert hdr == COLS17, f"Master-Header weicht ab: {hdr}"
    rows = []
    for v in it:
        if v is None or all(x in (None, "") for x in v):
            continue
        rows.append({c: ("" if x is None else str(x).strip()) for c, x in zip(COLS17, v)})
    return rows


def to_grid(v):
    """Koordinate auf das Geräteraster runden (Grad + Minuten mit 3 Nachkommastellen)."""
    x = float(v); s = -1 if x < 0 else 1; x = abs(x)
    deg = int(x); mins = round((x - deg) * 60, 3)
    if mins >= 60: deg += 1; mins -= 60
    return f"{s * (deg + mins / 60):.6f}"


def normalize(r, warn):
    """Werte so schreiben, wie das Gerät sie speichert (am ID-52E am 26.09. per Gegen-Export ermittelt),
    damit der Import ohne «some data was corrected» durchläuft."""
    if r["Latitude"] and r["Longitude"]:
        r["Latitude"], r["Longitude"] = to_grid(r["Latitude"]), to_grid(r["Longitude"])
    if r["TONE"] == "DTCS":   # Repeaterliste kennt nur OFF/TONE/TSQL
        warn.append(f"{r['Name']} {r['Frequency']}: DTCS {r['Repeater Tone']} nicht möglich → OFF (in Speicherkanal ablegen!)")
        r["TONE"], r["Repeater Tone"] = "OFF", "88.5Hz"
    if r["Mode"] == "FM" and r["Dup"] == "OFF":   # FM-Simplex: Gerät löscht das Rufzeichen
        r["Repeater Call Sign"] = ""
    return r


def check(rows, dev):
    err = []
    for r in rows:
        if len(r["Name"]) > 16: err.append(f"Name >16: {r['Name']}")
        if len(r["Sub Name"]) > 8: err.append(f"Sub Name >8: {r['Sub Name']}")
        if len(r["Group Name"]) > 16: err.append(f"Gruppenname >16: {r['Group Name']}")
        if len(r["Repeater Call Sign"]) > 8: err.append(f"Call >8: {r['Repeater Call Sign']}")
        if not all(32 <= ord(ch) < 127 for ch in "".join(r.values())): err.append(f"Nicht-ASCII: {r['Name']}")
    dup = [k for k, n in Counter((r["Group No"], r["Repeater Call Sign"], r["Mode"]) for r in rows
                                 if r["Mode"] in ("DV", "DD")).items() if n > 1]
    if dup: err.append(f"doppelte DV-Rufzeichen je Gruppe: {dup[:5]}")
    if len(rows) > dev["limit"]: err.append(f"{len(rows)} > Limit {dev['limit']}")
    if len({r['Group No'] for r in rows}) > 50: err.append("mehr als 50 Gruppen")
    return err


def main():
    master, outdir = sys.argv[1], sys.argv[2]
    datum = sys.argv[3] if len(sys.argv) > 3 else datetime.date.today().strftime("%Y%m%d")
    rows = read_master(master)
    warn = []
    rows = [normalize(r, warn) for r in rows]
    report = []
    for name, dev in DEVICES.items():
        variants = [("10", lambda r: True)]
        if name == "IC-9700":
            variants.append(("11", lambda r: r["Mode"] in ("DV", "DD")))
        for nn, flt in variants:
            sel = [dict(r) for r in rows if band(float(r["Frequency"])) in dev["bands"] and flt(r)]
            # Gruppen lückenlos neu nummerieren (Reihenfolge wie Master)
            order = OrderedDict()
            for r in sel:
                order.setdefault(r["Group No"], len(order) + 1)
            for r in sel:
                r["Group No"] = str(order[r["Group No"]])
            err = check(sel, dev)
            if err:
                print(f"FEHLER {name}:", *err, sep="\n  "); sys.exit(1)
            os.makedirs(os.path.join(outdir, name), exist_ok=True)
            path = os.path.join(outdir, name, f"Rpt{datum}_{nn}.csv")
            with open(path, "w", newline="", encoding="ascii") as f:
                w = csv.writer(f, delimiter=dev["sep"], lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
                w.writerow(dev["cols"])
                for r in sel:
                    out = dict(r)
                    if dev["dec"] == ",":
                        for c in NUM:
                            out[c] = out[c].replace(".", ",")
                    out.setdefault("Gateway IP Address", "")
                    w.writerow([out[c] for c in dev["cols"]])
            groups = Counter((r["Group No"], r["Group Name"]) for r in sel)
            report.append((name, nn, path, len(sel), groups))
    for w in warn:
        print("HINWEIS:", w)
    for name, nn, path, n, groups in report:
        print(f"{name:8} {os.path.basename(path)}  {n:5} Einträge, {len(groups)} Gruppen: " +
              ", ".join(f"{g}={c}" for (_, g), c in sorted(groups.items(), key=lambda x: int(x[0][0]))))


if __name__ == "__main__":
    main()
